"""8th-grade admission cut-offs from РУО София-град's min/max score spreadsheets.

After each 8th-grade round the РУО publishes "СПРАВКА за минималния и максималния бал по
паралелки": one row per school and class (паралелка) with the lowest and highest score
admitted in that round, overall and for boys and girls. Each справка covers the children
admitted in that round only, so later rounds are often lower.

Known files are listed in ``ruo_sofia_min_max_sources.json`` next to this module (file
names change every year and the post pages sit behind a Cloudflare check, but the files
themselves download directly). Layout quirks handled here, from the 2022–2026 files:
- the header is found by its text ("Код" row), not by position;
- 2023 adds a "Брой паралелки" column, 2026 round 2 has no "№" column;
- quota classes (прием по квоти) are marked by a coloured font (blue, green in 2022);
- class codes lose leading zeros up to 2025 (int 4486 vs "0814") and are 5 digits from 2026;
- a row of all zeros means nobody was admitted to that class in that round.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

SOURCES_FILE = Path(__file__).with_name("ruo_sofia_min_max_sources.json")
SOURCE = "ruo-sofia-grad"

_DATE_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4})")
ROUND_NAMES = {1: "Първо класиране", 2: "Второ класиране", 3: "Трето класиране", 4: "Четвърто класиране"}
# Наредба № 10/2016 amendments known to the rule files (frontend/src/admission/rules/bg/grade_8);
# a ranking is labelled with the latest one in force on its date. Earlier amendments are not
# yet researched, so older rankings carry no version.
RULE_VERSIONS = ("2025-09-26",)  # ДВ бр. 79/2025
# Plain text is black: theme colour 1 (up to 2025) or FF000000 (2026).
_PLAIN_RGB = {"FF000000", "00000000"}


@dataclass
class MinMaxRow:
    school_code: str
    school_name: str
    class_code: str
    class_name: str
    min_score: Optional[float]
    max_score: Optional[float]
    min_boys: Optional[float]
    min_girls: Optional[float]
    quota: bool


def load_sources() -> list[dict[str, Any]]:
    return json.loads(SOURCES_FILE.read_text(encoding="utf-8"))["files"]


def _text(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _score(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    try:
        number = float(str(value).replace(",", "."))
    except ValueError:
        return None
    return None if number == 0 else number


def class_code(value: Any, year: int) -> str:
    raw = str(value).strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    return raw.zfill(5 if year >= 2026 else 4)


def _is_coloured(cell) -> bool:
    color = getattr(cell.font, "color", None)
    if color is None:
        return False
    if color.type == "rgb":
        return isinstance(color.rgb, str) and color.rgb.upper() not in _PLAIN_RGB
    if color.type == "indexed":
        # 8 is black and 64 the system text colour in the legacy palette.
        return color.indexed not in (None, 8, 64)
    if color.type == "theme":
        # Theme 1 is the default text colour; 2023 marks quota rows with theme 4 (accent).
        return color.theme not in (None, 1)
    return False


def report_date(sheet) -> Optional[date]:
    for row in sheet.iter_rows(min_row=1, max_row=4, values_only=True):
        for value in row:
            if isinstance(value, str) and "справка" in value.lower():
                match = _DATE_RE.search(value)
                if match:
                    day, month, year = map(int, match.groups())
                    return date(year, month, day)
    return None


def _header_columns(sheet) -> tuple[int, dict[str, int]]:
    """(first data row, column index by field) from the two header rows."""
    for row_index in range(1, 15):
        values = [_text(cell.value).lower() for cell in sheet[row_index]]
        if values.count("код") >= 2:
            top = [_text(cell.value).lower() for cell in sheet[row_index - 1]]
            codes = [i for i, v in enumerate(values) if v == "код"]
            totals = [i for i, v in enumerate(values) if v == "общо"]
            # Fill merged top-header labels rightwards to tell min from max.
            spans, current = [], ""
            for v in top:
                current = v or current
                spans.append(current)
            min_total = next(i for i in totals if "мин" in spans[i])
            max_total = next(i for i in totals if "макс" in spans[i])
            columns = {
                "school_code": codes[0],
                "school_name": codes[0] + 1,
                "class_code": codes[1],
                "class_name": codes[1] + 1,
                "min_score": min_total,
                "min_boys": min_total + 1,
                "min_girls": min_total + 2,
                "max_score": max_total,
            }
            return row_index + 1, columns
    raise ValueError("no header row with two 'Код' cells")


def parse_workbook(content: bytes, year: int) -> tuple[Optional[date], list[MinMaxRow]]:
    sheet = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    first_row, columns = _header_columns(sheet)
    rows = []
    for cells in sheet.iter_rows(min_row=first_row):
        school_code = _text(cells[columns["school_code"]].value)
        if school_code.endswith(".0"):
            school_code = school_code[:-2]
        if not school_code.isdigit():
            continue  # blank spacer or a footer note
        values = {key: cells[index].value for key, index in columns.items()}
        rows.append(
            MinMaxRow(
                school_code=school_code,
                school_name=_text(values["school_name"]),
                class_code=class_code(values["class_code"], year),
                class_name=_text(values["class_name"]),
                min_score=_score(values["min_score"]),
                max_score=_score(values["max_score"]),
                min_boys=_score(values["min_boys"]),
                min_girls=_score(values["min_girls"]),
                quota=_is_coloured(cells[columns["class_name"]]),
            )
        )
    return report_date(sheet), rows


def rule_set_for(year: int, classified_on: date) -> str:
    versions = [v for v in RULE_VERSIONS if v <= classified_on.isoformat()]
    suffix = f"@{versions[-1]}" if versions else ""
    return f"bg/grade_8/{year}-{year + 1}{suffix}"


def archive_rows(source: dict[str, Any], classified_on: Optional[date], rows: list[MinMaxRow], retrieved_at: str) -> list[dict[str, Any]]:
    """Archive rows for one справка; classes nobody was admitted to are left out."""
    year = int(source["year"])
    campaign = f"{year}/{year + 1}"
    classified_on = classified_on or date.fromisoformat(source["date"])
    out = []
    for row in rows:
        if row.min_score is None and row.max_score is None:
            continue
        item = {
            "country": "bg",
            "city": "sofia",
            "stage": "grade_8",
            "campaign": campaign,
            "rule_set": rule_set_for(year, classified_on),
            "round_id": int(source["round"]),
            "classified_on": classified_on.isoformat(),
            "round_name": ROUND_NAMES.get(int(source["round"])),
            "round_kind": "general",
            "source": SOURCE,
            "source_id": row.school_code,
            "institution_name": row.school_name,
            "class_code": row.class_code,
            "class_name": row.class_name,
            "queue": "general",
            "quota": row.quota,
            "min_score": row.min_score,
            "max_score": row.max_score,
            "source_url": source["url"],
            "retrieved_at": retrieved_at,
        }
        if row.quota:
            item["min_score_boys"] = row.min_boys
            item["min_score_girls"] = row.min_girls
        out.append(item)
    return out
