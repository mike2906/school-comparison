"""РУО София-град 8th-grade min/max spreadsheets (synthetic workbooks in the real layouts)."""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from app.scrapers.sources.bg.ruo_sofia_min_max import (
    archive_rows,
    class_code,
    load_sources,
    parse_workbook,
    rule_set_for,
)

BLUE = Font(color="FF0070C0")
TITLE = "СПРАВКА\nза минималния и максималния бал по паралелки\nв РУО СОФИЯ-ГРАД\nкъм дата 21.07.2025 г."


def _workbook(*, numbered=True, count_column=False, quota_font=BLUE, rows=None):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = TITLE
    ws["B2"] = "Забележка: Оцветените в синьо паралелки са с прием по квоти!!!"
    lead = ["№"] if numbered else []
    extra = ["Брой паралелки"] if count_column else []
    top = lead + ["Училище", None, "Паралелка", None] + extra + ["Минимален бал", None, None, "Максимален бал", None, None]
    bottom = ([None] if numbered else []) + ["Код", "Име", "Код", "Име"] + ([None] if count_column else []) + ["Общо", "Мъже", "Жени", "Общо", "Мъже", "Жени"]
    ws.append(top)  # row 3
    ws.append(bottom)  # row 4
    ws.append([])  # row 5 blank
    rows = rows or [
        (2216306, "СМГ", 3909, "Математически", 492, 492, 0, 493, 493, 0, False),
        (2216164, "164. ГПИЕ", "0814", "Чужди езици - ИЕ интензивно, АЕ", 475, 475, 483.5, 484, 480, 484, True),
        (2216081, "81. СУ", 4486, "Хуманитарни науки", 0, 0, 0, 0, 0, 0, False),
    ]
    for i, (school, name, code, cls, *scores, quota) in enumerate(rows, 1):
        values = ([i] if numbered else []) + [school, name, code, cls] + ([1] if count_column else []) + scores
        ws.append(values)
        if quota:
            for cell in ws[ws.max_row]:
                cell.font = quota_font
    ws.append([None, "Забележка: За паралелките с профил Физическо възпитание и спорт се кандидатства в училището"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parses_rows_date_and_quota():
    classified_on, rows = parse_workbook(_workbook(), 2025)
    assert classified_on == date(2025, 7, 21)
    assert [r.school_code for r in rows] == ["2216306", "2216164", "2216081"]
    smg, quota, empty = rows
    assert (smg.class_code, smg.min_score, smg.max_score, smg.quota) == ("3909", 492.0, 493.0, False)
    assert (quota.class_code, quota.quota, quota.min_boys, quota.min_girls) == ("0814", True, 475.0, 483.5)
    assert empty.min_score is None and empty.max_score is None


def test_layout_variants_found_by_header_text():
    _, without_number = parse_workbook(_workbook(numbered=False), 2026)
    assert without_number[0].class_code == "03909" and without_number[0].min_score == 492.0
    _, with_count = parse_workbook(_workbook(count_column=True), 2023)
    assert with_count[1].min_score == 475.0 and with_count[1].max_score == 484.0
    _, green = parse_workbook(_workbook(quota_font=Font(color="FF00B050")), 2022)
    assert [r.quota for r in green] == [False, True, False]


def test_class_codes_pad_by_year():
    assert class_code(814, 2025) == "0814"
    assert class_code("0814", 2025) == "0814"
    assert class_code("03909", 2026) == "03909"
    assert class_code(4486.0, 2026) == "04486"


def test_archive_rows_skip_classes_nobody_entered():
    classified_on, rows = parse_workbook(_workbook(), 2025)
    source = {"year": 2025, "round": 2, "date": "2025-07-21", "url": "https://example.test/f.xlsx"}
    out = archive_rows(source, classified_on, rows, "t")
    assert len(out) == 2
    smg, quota = out
    assert (smg["stage"], smg["campaign"], smg["round_id"], smg["round_name"]) == ("grade_8", "2025/2026", 2, "Второ класиране")
    assert smg["rule_set"] == "bg/grade_8/2025-2026"  # July 2025, before ДВ 79/2025
    assert (smg["source_id"], smg["queue"], smg["quota"]) == ("2216306", "general", False)
    assert "min_score_boys" not in smg
    assert (quota["quota"], quota["min_score_boys"], quota["min_score_girls"]) == (True, 475.0, 483.5)


def test_rule_set_names_the_version_in_force():
    assert rule_set_for(2025, date(2025, 7, 21)) == "bg/grade_8/2025-2026"
    assert rule_set_for(2026, date(2026, 7, 14)) == "bg/grade_8/2026-2027@2025-09-26"


def test_source_list_is_well_formed():
    sources = load_sources()
    assert len({s["url"] for s in sources}) == len(sources)
    for source in sources:
        assert source["round"] in (1, 2, 3, 4)
        assert date.fromisoformat(source["date"]).year == source["year"]
        assert source["url"].endswith(".xlsx")


def test_indexed_palette_colour_marks_quota():
    from openpyxl.styles.colors import Color

    _, rows = parse_workbook(_workbook(quota_font=Font(color=Color(indexed=12))), 2025)
    assert [r.quota for r in rows] == [False, True, False]
